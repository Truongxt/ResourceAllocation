from pathlib import Path
from copy import deepcopy
from textwrap import wrap

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\admin\Downloads\218_110_Truong_Tien_TaiLieu (1).docx")
OUTPUT = ROOT / "docs" / "218_110_Truong_Tien_TaiLieu_HoanThien.docx"
ASSETS = ROOT / "docs" / "report_assets"
ASSETS.mkdir(exist_ok=True)


def font(size=34, bold=False):
    candidates = [
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\timesbd.ttf" if bold else r"C:\Windows\Fonts\times.ttf",
    ]
    for name in candidates:
        if Path(name).exists():
            return ImageFont.truetype(name, size)
    return ImageFont.load_default()


def rounded_box(draw, xy, title, lines, fill, outline=(28, 63, 104), title_fill=(255, 255, 255)):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=24, fill=fill, outline=outline, width=4)
    draw.rounded_rectangle((x1, y1, x2, y1 + 58), radius=24, fill=outline)
    draw.rectangle((x1, y1 + 30, x2, y1 + 58), fill=outline)
    draw.text((x1 + 20, y1 + 12), title, font=font(29, True), fill=title_fill)
    yy = y1 + 78
    for line in lines:
        draw.text((x1 + 22, yy), line, font=font(24), fill=(25, 35, 48))
        yy += 34


def arrow(draw, start, end, color=(42, 86, 135), width=5):
    draw.line((start, end), fill=color, width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) > abs(ey - sy):
        sign = 1 if ex > sx else -1
        pts = [(ex, ey), (ex - 18 * sign, ey - 11), (ex - 18 * sign, ey + 11)]
    else:
        sign = 1 if ey > sy else -1
        pts = [(ex, ey), (ex - 11, ey - 18 * sign), (ex + 11, ey - 18 * sign)]
    draw.polygon(pts, fill=color)


def make_diagrams():
    # Architecture
    img = Image.new("RGB", (1800, 1120), "white")
    d = ImageDraw.Draw(img)
    d.text((900, 35), "KIẾN TRÚC TỔNG THỂ HỆ THỐNG RAO", anchor="ma", font=font(42, True), fill=(19, 54, 93))
    rounded_box(d, (160, 110, 1640, 300), "Tầng trình bày", ["Web: React 18, Ant Design 6, React Router, Axios", "Mobile bổ trợ: React Native / Expo", "Giao diện sáng–tối, Việt–Anh, dashboard, Kanban, Gantt, báo cáo"], (232, 244, 255))
    rounded_box(d, (160, 390, 1640, 655), "Tầng dịch vụ ứng dụng", ["Node.js + Express: 12 nhóm route, controller, validation và phân quyền", "REST API cho nghiệp vụ; Socket.IO cho thông báo thời gian thực", "Bảo mật: JWT access/refresh, bcrypt, Helmet, CORS, rate limiting", "Khối tối ưu hóa: Genetic Algorithm, CSP và Hybrid"], (235, 249, 240))
    rounded_box(d, (160, 745, 1640, 1010), "Tầng dữ liệu", ["MongoDB 7 + Mongoose 8", "12 mô hình: User, RefreshToken, Project, Task, TaskGroup, RecurringTask,", "Resource, Department, OptimizationResult, Notification, ActivityLog, CompanySetting"], (255, 245, 224))
    arrow(d, (900, 302), (900, 388))
    arrow(d, (900, 657), (900, 743))
    img.save(ASSETS / "architecture.png", quality=95)

    # Data model
    img = Image.new("RGB", (1800, 1180), "white")
    d = ImageDraw.Draw(img)
    d.text((900, 35), "MÔ HÌNH DỮ LIỆU KHÁI QUÁT", anchor="ma", font=font(42, True), fill=(19, 54, 93))
    boxes = {
        "User": (70, 130, 410, 280), "Project": (730, 120, 1070, 270), "Task": (1390, 130, 1730, 280),
        "Resource": (70, 500, 410, 650), "Department": (70, 840, 410, 990),
        "TaskGroup": (730, 430, 1070, 580), "RecurringTask": (730, 760, 1070, 910),
        "OptimizationResult": (1390, 500, 1730, 650), "Notification": (1390, 840, 1730, 990),
        "RefreshToken": (500, 955, 840, 1105), "ActivityLog": (960, 955, 1300, 1105),
    }
    colors = [(226, 240, 255), (232, 248, 237), (255, 242, 220)]
    for i, (name, xy) in enumerate(boxes.items()):
        rounded_box(d, xy, name, [], colors[i % 3])
    for a, b in [
        ((410, 205), (730, 195)), ((1070, 195), (1390, 205)), ((240, 280), (240, 500)),
        ((410, 575), (730, 505)), ((1070, 505), (1390, 575)), ((240, 650), (240, 840)),
        ((900, 580), (900, 760)), ((1560, 650), (1560, 840)), ((410, 225), (670, 955)),
        ((1530, 280), (1130, 955)),
    ]:
        arrow(d, a, b, width=4)
    d.text((900, 320), "Dự án chứa công việc; công việc gắn người dùng và nhóm", anchor="ma", font=font(25), fill=(70, 82, 96))
    d.text((900, 680), "Kết quả tối ưu lưu phương án, chỉ số và trạng thái áp dụng", anchor="ma", font=font(25), fill=(70, 82, 96))
    img.save(ASSETS / "data_model.png", quality=95)

    # Optimization flow
    img = Image.new("RGB", (1800, 1140), "white")
    d = ImageDraw.Draw(img)
    d.text((900, 35), "LUỒNG TỐI ƯU HÓA PHÂN BỔ NHÂN SỰ", anchor="ma", font=font(42, True), fill=(19, 54, 93))
    steps = [
        ("1. Chuẩn bị dữ liệu", ["Task mở, kỹ năng yêu cầu, effort", "Resource hoạt động, năng lực, lịch nghỉ"]),
        ("2. Kiểm tra khả thi", ["Kỹ năng ≥ ngưỡng", "Không giao lịch nghỉ; capacity hợp lệ"]),
        ("3. Chọn bộ giải", ["GA: tìm kiếm đa mục tiêu", "CSP: ràng buộc cứng", "Hybrid: CSP thu hẹp miền → GA"]),
        ("4. Đánh giá", ["Fitness, skill match, độ lệch tải", "utilization, quá tải, chi phí, thời gian"]),
        ("5. Phê duyệt", ["Lưu lịch sử và so sánh phương án", "PM/Admin áp dụng hoặc hoàn tác"]),
    ]
    ys = [110, 310, 510, 710, 910]
    for i, (title, lines) in enumerate(steps):
        rounded_box(d, (260, ys[i], 1540, ys[i] + 145), title, lines, [(232, 244, 255), (235, 249, 240), (255, 245, 224)][i % 3])
        if i < len(steps) - 1:
            arrow(d, (900, ys[i] + 147), (900, ys[i + 1] - 2))
    img.save(ASSETS / "optimization_flow.png", quality=95)

    # Testing result
    img = Image.new("RGB", (1800, 930), "white")
    d = ImageDraw.Draw(img)
    d.text((900, 35), "KẾT QUẢ KIỂM CHỨNG TẠI THỜI ĐIỂM LẬP BÁO CÁO", anchor="ma", font=font(40, True), fill=(19, 54, 93))
    rounded_box(d, (100, 140, 830, 410), "Build production", ["4.218 mô-đun được chuyển đổi", "Build thành công trong 25,58 giây", "Có cảnh báo chunk chính > 500 kB"], (232, 248, 238))
    rounded_box(d, (970, 140, 1700, 410), "Kiểm thử logic client", ["Gantt/đường găng: 33/33 đạt", "Đa ngôn ngữ & enum: 12/12 đạt", "Tổng kiểm thử logic: 45/45 đạt"], (232, 244, 255))
    rounded_box(d, (100, 520, 830, 790), "Kiểm thử component", ["27/31 ca đạt; 4 ca chưa đạt", "Nguyên nhân chính: kỳ vọng test lệch UI", "Cần xử lý trước nghiệm thu cuối"], (255, 242, 220), outline=(180, 105, 25))
    rounded_box(d, (970, 520, 1700, 790), "Kiểm thử backend", ["Có bộ test API, Socket, CSP, Hybrid", "Phần tích hợp cần MongoDB test riêng", "Không tuyên bố đạt nếu chưa chạy đủ môi trường"], (245, 240, 252), outline=(102, 73, 145))
    img.save(ASSETS / "verification.png", quality=95)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_text(cell, value, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=11):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(str(value))
    r.bold = bold
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True, WD_ALIGN_PARAGRAPH.CENTER, 10.5)
        shade_cell(table.rows[0].cells[i], "D9EAF7")
    set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, WD_ALIGN_PARAGRAPH.LEFT, 10.5)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    return table


def add_field(paragraph, instruction, placeholder="Nhấn Ctrl+A, F9 để cập nhật trường"):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    sep = OxmlElement('w:fldChar')
    sep.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = placeholder
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for node in (begin, instr, sep, text, end):
        run._r.append(node)


def insert_frontmatter(doc):
    anchor = next(p for p in doc.paragraphs if p.text.strip() == "LỜI CẢM ƠN")
    blocks = []
    for title, field in [
        ("MỤC LỤC", 'TOC \\o "1-3" \\h \\z \\u'),
        ("DANH MỤC CÁC HÌNH ẢNH", 'TOC \\h \\z \\t "Figure Caption,1"'),
        ("DANH MỤC CÁC BẢNG BIỂU", 'TOC \\h \\z \\t "Table Caption,1"'),
    ]:
        p = OxmlElement('w:p')
        p_pr = OxmlElement('w:pPr')
        p_style = OxmlElement('w:pStyle')
        p_style.set(qn('w:val'), 'Title')
        p_pr.append(p_style)
        p.append(p_pr)
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = title
        r.append(t)
        p.append(r)
        blocks.append(p)
        fp = OxmlElement('w:p')
        fr = OxmlElement('w:r')
        begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
        instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = field
        sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
        txt = OxmlElement('w:t'); txt.text = "Mục này sẽ tự cập nhật khi mở bằng Microsoft Word."
        end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
        for node in (begin, instr, sep, txt, end): fr.append(node)
        fp.append(fr)
        blocks.append(fp)
        brp = OxmlElement('w:p')
        brr = OxmlElement('w:r')
        br = OxmlElement('w:br'); br.set(qn('w:type'), 'page')
        brr.append(br); brp.append(brr); blocks.append(brp)
    for block in blocks:
        anchor._p.addprevious(block)


def add_style(doc, name, base="Caption", size=11, italic=True):
    if name in doc.styles:
        return doc.styles[name]
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles[base]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.italic = italic
    return style


def normalize_styles(doc):
    add_style(doc, "Figure Caption")
    add_style(doc, "Table Caption")
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(1)
    for name, size, color in [("Heading 1", 16, RGBColor(0, 0, 0)), ("Heading 2", 14, RGBColor(0, 0, 0)), ("Heading 3", 13, RGBColor(0, 0, 0))]:
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.keep_with_next = True


def p(doc, text="", bold_lead=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=False):
    para = doc.add_paragraph()
    para.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r = para.add_run(bold_lead); r.bold = True
        para.add_run(text[len(bold_lead):])
    else:
        r = para.add_run(text); r.italic = italic
    return para


def bullet(doc, text, level=0):
    para = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.add_run(text)
    return para


def numbered(doc, text):
    para = doc.add_paragraph(style="List Number")
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.add_run(text)
    return para


def caption(doc, text, kind="figure"):
    para = doc.add_paragraph(style="Figure Caption" if kind == "figure" else "Table Caption")
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run(text)
    return para


def figure(doc, image_name, caption_text, width=6.2):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Cm(0)
    para.add_run().add_picture(str(ASSETS / image_name), width=Inches(width))
    caption(doc, caption_text, "figure")


def heading(doc, text, level):
    para = doc.add_paragraph(text, style=f"Heading {level}")
    # The university template numbers chapter headings through a local numPr
    # rather than through the Heading 1 style itself. Reapply that numbering
    # to generated chapter titles so they continue after Chapter 1.
    if level == 1 and text.startswith(":"):
        p_pr = para._p.get_or_add_pPr()
        num_pr = OxmlElement('w:numPr')
        ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), '0')
        num_id = OxmlElement('w:numId'); num_id.set(qn('w:val'), '1')
        num_pr.append(ilvl); num_pr.append(num_id)
        p_pr.append(num_pr)
    return para


def update_existing(doc):
    replacements = {
        "– React 18 kết hợp công cụ đóng gói Vite 5 và thư viện giao diện Ant Design 6 — dùng để xây dựng toàn bộ giao diện phía trình duyệt, gồm 12 trang chức năng.":
            "– React 18 kết hợp Vite 5 và Ant Design 6 để xây dựng giao diện web theo mô hình SPA; các trang được tách theo route nhằm giảm tải ban đầu.",
        "– Không xây dựng các tính năng cộng tác trong công việc như bình luận, đính kèm tệp và nhắc tên người dùng.":
            "– Hệ thống đã hỗ trợ bình luận, checklist, người theo dõi và công việc con; chưa triển khai kho tệp đính kèm dung lượng lớn và nhắc tên nâng cao.",
        "Các nhóm tính năng chính của hệ thống gồm: xác thực và phân quyền theo vai trò; quản lý dự án và thành viên dự án; quản lý công việc với bảng Kanban kéo thả, quan hệ phụ thuộc và kỹ năng yêu cầu; quản lý nhân sự với ma trận kỹ năng, năng lực và lịch nghỉ; module tối ưu hóa phân bổ nhân sự với ba thuật toán, lưu lịch sử và so sánh phương án; biểu đồ Gantt đa dự án có đường tới hạn; báo cáo và phân tích mức sử dụng nguồn lực kèm cảnh báo quá tải; thông báo thời gian thực; và nhật ký hoạt động hệ thống.":
            "Các nhóm tính năng chính gồm: xác thực và phân quyền theo vai trò/quyền ứng dụng; quản lý dự án, thành viên và nhóm công việc; quản lý công việc bằng bảng Kanban, lịch, Gantt, phụ thuộc, công việc con, bình luận, checklist và người theo dõi; quản lý nhân sự với ma trận kỹ năng, năng lực và lịch nghỉ; công việc lặp; module tối ưu hóa GA/CSP/Hybrid có benchmark, lịch sử, so sánh, áp dụng và hoàn tác; báo cáo, thông báo thời gian thực, nhật ký hoạt động và cấu hình công ty. Phiên bản di động đóng vai trò kênh truy cập bổ trợ cho các nghiệp vụ chính.",
    }
    for para in doc.paragraphs:
        if para.text in replacements:
            para.text = replacements[para.text]
    # Restyle the existing table caption so it appears in the generated list.
    for para in doc.paragraphs:
        if para.text.strip().startswith("Bảng 1-1"):
            para.style = doc.styles["Table Caption"]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_chapter_2(doc):
    heading(doc, ": CƠ SỞ LÝ THUYẾT", 1)
    p(doc, "Chương này trình bày các khái niệm và phương pháp được sử dụng để hình thành mô hình nghiệp vụ, kiến trúc phần mềm và bộ máy tối ưu hóa của hệ thống. Nội dung tập trung vào quản lý đa dự án, mô hình hóa năng lực nhân sự, bài toán phân bổ có ràng buộc và ba phương pháp giải đã được hiện thực.")

    heading(doc, "Quản lý luồng công việc đa dự án", 2)
    p(doc, "Quản lý đa dự án là hoạt động lập kế hoạch và điều phối đồng thời nhiều dự án sử dụng chung một tập nguồn lực. Khác với quản lý một dự án độc lập, quyết định giao việc ở dự án này có thể làm giảm năng lực sẵn có của dự án khác. Vì vậy hệ thống cần cung cấp một góc nhìn hợp nhất về thời gian, khối lượng và trách nhiệm của từng nhân sự.")
    p(doc, "Luồng công việc được mô hình hóa bằng tập các trạng thái cố định gồm Cần làm, Đang làm, Đánh giá, Hoàn thành và Bị chặn. Quan hệ phụ thuộc giữa các công việc tạo thành đồ thị có hướng; chu trình trong đồ thị là dữ liệu không hợp lệ. Biểu đồ Gantt dùng các mốc bắt đầu/kết thúc và quan hệ phụ thuộc để biểu diễn tiến độ; phương pháp đường găng (Critical Path Method – CPM) xác định chuỗi công việc có tổng thời lượng dài nhất và độ dự trữ bằng không.")

    heading(doc, "Mô hình năng lực và mức sử dụng nguồn lực", 2)
    p(doc, "Mỗi nhân sự được mô tả bởi ma trận kỹ năng, năng lực tối đa theo tuần, hệ số toàn thời gian (FTE), trạng thái sẵn sàng, lịch nghỉ và đơn giá giờ. Mỗi công việc chứa số giờ ước tính, khoảng thời gian thực hiện, mức ưu tiên và danh sách kỹ năng yêu cầu. Cặp dữ liệu này cho phép hệ thống chấm điểm độ phù hợp và dự báo quá tải trước khi quyết định phân công.")
    p(doc, "Mức sử dụng năng lực của nhân sự r trong một kỳ được tính theo công thức:", align=WD_ALIGN_PARAGRAPH.LEFT)
    eq = p(doc, "Utilization(r) = Workload(r) / (MaxCapacity(r) × FTE(r)) × 100%", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
    eq.paragraph_format.first_line_indent = Cm(0)
    p(doc, "Khi chỉ số vượt 100%, nhân sự bị xem là quá tải. Tuy nhiên chỉ số tải chưa đủ để kết luận phân công tốt: một phương án cân bằng nhưng giao sai chuyên môn vẫn có thể gây rủi ro chất lượng. Do đó hệ thống kết hợp cả tải, kỹ năng, chi phí và số người quá tải trong hàm mục tiêu.")

    heading(doc, "Bài toán phân bổ nhân sự", 2)
    p(doc, "Cho tập công việc T = {t₁, …, tₙ} và tập nhân sự R = {r₁, …, rₘ}. Biến quyết định xᵢ nhận giá trị là nhân sự được giao cho công việc tᵢ. Một phương án là véc-tơ X = [x₁, …, xₙ]. Không gian tìm kiếm có kích thước mⁿ, tăng theo hàm mũ; vì vậy duyệt toàn bộ chỉ phù hợp với dữ liệu rất nhỏ.")
    p(doc, "Điểm khớp kỹ năng có trọng số giữa công việc t và nhân sự r được tính như sau:", align=WD_ALIGN_PARAGRAPH.LEFT)
    eq = p(doc, "SkillMatch(t,r) = Σ[wₛ × min(Level(r,s), Required(t,s))] / Σ[wₛ × Required(t,s)]", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
    eq.paragraph_format.first_line_indent = Cm(0)
    p(doc, "Thang kỹ năng được chuẩn hóa từ 1 đến 4. Việc dùng hàm min bảo đảm một nhân sự vượt yêu cầu không tạo điểm lớn hơn 100%, còn trọng số cho phép kỹ năng cốt lõi ảnh hưởng mạnh hơn kỹ năng phụ.")

    heading(doc, "Giải thuật di truyền", 2)
    p(doc, "Giải thuật di truyền (Genetic Algorithm – GA) là phương pháp tìm kiếm gần đúng mô phỏng quá trình tiến hóa. Mỗi nhiễm sắc thể là một phương án phân công; gene thứ i chứa chỉ số nhân sự thực hiện công việc i. Quần thể ban đầu được sinh ngẫu nhiên, sau đó lặp qua các bước đánh giá, chọn lọc, lai ghép, đột biến và giữ lại cá thể ưu tú.")
    p(doc, "Hàm thích nghi tổng hợp bốn mục tiêu đã chuẩn hóa:", align=WD_ALIGN_PARAGRAPH.LEFT)
    eq = p(doc, "F = 0,30·f_workload + 0,35·f_skill + 0,15·f_cost + 0,20·f_overallocation", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
    eq.paragraph_format.first_line_indent = Cm(0)
    add_table(doc, ["Thành phần", "Ý nghĩa", "Trọng số mặc định"], [
        ("f_workload", "Ưu tiên độ lệch tải thấp", "0,30"),
        ("f_skill", "Ưu tiên kỹ năng phù hợp", "0,35"),
        ("f_cost", "Ưu tiên tổng chi phí thấp", "0,15"),
        ("f_overallocation", "Phạt số nhân sự quá tải", "0,20"),
    ], [3.5, 9.0, 3.5])
    caption(doc, "Bảng 2-1 Các thành phần của hàm thích nghi", "table")
    p(doc, "Hệ thống sử dụng chọn lọc theo giải đấu k = 5, lai ghép đồng nhất, đột biến gán lại ngẫu nhiên và elitism 5%. Thuật toán dừng khi đạt số thế hệ tối đa, fitness vượt ngưỡng mục tiêu hoặc không cải thiện trong 50 thế hệ liên tiếp. GA có ưu điểm xử lý tốt mục tiêu mềm và dữ liệu lớn, nhưng không bảo đảm tối ưu toàn cục và cần hiệu chỉnh tham số.")

    heading(doc, "Bài toán thỏa mãn ràng buộc", 2)
    p(doc, "CSP biểu diễn mỗi công việc như một biến, miền của biến là tập nhân sự có thể nhận việc, và tập ràng buộc mô tả điều kiện hợp lệ. Bộ giải trong project dùng quay lui kết hợp MRV, LCV, node consistency và AC-3. MRV chọn trước biến có ít lựa chọn nhất; LCV thử trước nhân sự còn nhiều năng lực để hạn chế khóa miền của các biến sau.")
    add_table(doc, ["Mã", "Ràng buộc", "Cách kiểm tra"], [
        ("H1", "Năng lực", "Tổng effort được giao không vượt capacity × FTE"),
        ("H2", "Kỹ năng", "Điểm khớp tổng hợp đạt ngưỡng mặc định 0,5"),
        ("H3", "Khả dụng", "Không giao vào thời kỳ nhân sự unavailable hoặc nghỉ"),
        ("H4", "Phụ thuộc", "Hai công việc phụ thuộc và chồng lịch không giao cùng người"),
    ], [2.0, 4.0, 10.0])
    caption(doc, "Bảng 2-2 Các ràng buộc cứng của CSP", "table")
    p(doc, "AC-3 duy trì tính nhất quán cung trên các cặp biến có ràng buộc nhị phân. Với quan hệ khác nhau (≠), một giá trị chỉ bị loại khi biến kề đã bị ép về đúng giá trị đó. Cơ chế này phát hiện sớm một số trường hợp vô nghiệm, nhưng không thay thế được ràng buộc all-different toàn cục.")

    heading(doc, "Phương pháp lai CSP–GA", 2)
    p(doc, "Phương pháp lai kết hợp tính chặt chẽ của CSP và khả năng tối ưu đa mục tiêu của GA. CSP lọc miền theo kỹ năng, lịch nghỉ và năng lực; GA chỉ sinh gene trong miền khả thi đó. Nhờ vậy quần thể không tiêu tốn nhiều thế hệ để học lại các điều kiện chắc chắn sai. Nếu một công việc có miền rỗng, hệ thống mở lại toàn bộ nhân sự cho riêng công việc đó và ghi cảnh báo tasksReopened, tránh trả kết quả có vẻ hợp lệ nhưng che giấu việc nới ràng buộc.")
    figure(doc, "optimization_flow.png", "Hình 2-1 Luồng tối ưu hóa phân bổ nhân sự")

    heading(doc, "Nền tảng công nghệ và nguyên tắc bảo mật", 2)
    p(doc, "Frontend sử dụng React để xây dựng giao diện theo thành phần, React Router cho điều hướng một trang và Ant Design cho hệ thống thành phần giao diện. Backend sử dụng Express để tổ chức REST API, Mongoose ánh xạ đối tượng sang MongoDB và Socket.IO truyền sự kiện thời gian thực. Ứng dụng di động bổ trợ dùng React Native với Expo.")
    p(doc, "Xác thực sử dụng access token ngắn hạn lưu trong bộ nhớ phía client và refresh token gửi bằng cookie httpOnly. Refresh token chỉ lưu bản băm trong cơ sở dữ liệu, được xoay vòng và có thể thu hồi theo phiên. Mật khẩu được băm bằng bcrypt; Helmet thiết lập các header bảo mật; CORS giới hạn nguồn; rate limiter hạn chế thử đăng nhập sai và lưu lượng API. Dữ liệu đầu vào được kiểm tra ở route bằng express-validator và chuẩn hóa lỗi trước khi trả về client.")


def build_chapter_3(doc):
    heading(doc, ": PHÂN TÍCH – THIẾT KẾ", 1)
    heading(doc, "Phân tích tác nhân và yêu cầu", 2)
    p(doc, "Hệ thống phục vụ ba vai trò cơ sở: Quản trị viên, Quản lý dự án và Thành viên. Ngoài vai trò, phiên bản hiện tại bổ sung quyền ứng dụng, quyền theo dự án, quyền đặc biệt và quan hệ quản lý trực tiếp. Cách tiếp cận này giúp tránh việc gắn mọi quyết định truy cập vào một nhãn vai trò duy nhất.")
    add_table(doc, ["Tác nhân", "Nghiệp vụ trọng tâm", "Ranh giới quyền"], [
        ("Quản trị viên", "Tài khoản, phòng ban, cấu hình, nhật ký, toàn bộ nghiệp vụ", "Có quyền quản trị toàn cục; các thao tác nhạy cảm được ghi log"),
        ("Quản lý dự án", "Dự án, thành viên, công việc, nhân sự, chạy và áp dụng tối ưu", "Chỉ thao tác trong phạm vi được giao và quyền ứng dụng"),
        ("Thành viên", "Theo dõi công việc, cập nhật tiến độ, lịch nghỉ cá nhân, thông báo", "Không thay đổi phạm vi công việc hoặc phân công của người khác"),
    ], [3.0, 7.0, 6.0])
    caption(doc, "Bảng 3-1 Tác nhân và ranh giới quyền", "table")

    heading(doc, "Các luồng nghiệp vụ chính", 2)
    heading(doc, "Luồng quản lý dự án và công việc", 3)
    numbered(doc, "Quản lý dự án tạo dự án, xác định thời gian, độ ưu tiên và thành viên cùng tỷ lệ tham gia.")
    numbered(doc, "Công việc được tạo trong dự án/nhóm, kèm effort, ngày, kỹ năng, phụ thuộc, người thực hiện và trạng thái.")
    numbered(doc, "Thành viên cập nhật trạng thái, tiến độ, giờ thực tế, kết quả, checklist và bình luận trong phạm vi được phép.")
    numbered(doc, "Hệ thống tự tính lại tiến độ dự án, tải nhân sự, thông báo và nhật ký hoạt động.")
    numbered(doc, "Người quản lý theo dõi trên bảng, Kanban, lịch, Gantt, dashboard và báo cáo.")
    heading(doc, "Luồng tối ưu hóa", 3)
    numbered(doc, "Chọn phạm vi toàn hệ thống hoặc một dự án và cấu hình bộ giải.")
    numbered(doc, "Server lấy các công việc chưa hoàn thành và nhân sự đang hoạt động, sau đó kiểm tra tính sẵn sàng của dữ liệu.")
    numbered(doc, "Bộ giải GA, CSP hoặc Hybrid tạo phương án và các chỉ số định lượng.")
    numbered(doc, "Kết quả được lưu để xem lịch sử, benchmark và so sánh nhiều phương án.")
    numbered(doc, "Quản lý dự án hoặc quản trị viên phê duyệt để áp dụng; hệ thống lưu trạng thái trước khi áp dụng nhằm hỗ trợ hoàn tác.")

    heading(doc, "Kiến trúc hệ thống", 2)
    p(doc, "Hệ thống áp dụng kiến trúc client–server ba tầng. Tầng trình bày chịu trách nhiệm hiển thị và tương tác. Tầng dịch vụ ứng dụng cung cấp API, kiểm tra quyền, điều phối nghiệp vụ và chạy thuật toán. Tầng dữ liệu lưu trữ trạng thái bền vững. Kênh Socket.IO hoạt động song song với REST API để chuyển thông báo theo phòng user:<id>.")
    figure(doc, "architecture.png", "Hình 3-1 Kiến trúc tổng thể hệ thống")
    p(doc, "Trong backend, route định nghĩa URL, validation và middleware quyền; controller xử lý nghiệp vụ và gọi trực tiếp Mongoose model. Thư mục services dành cho các mối quan tâm dùng chung như Socket.IO, email, refresh token và nhật ký. Cách tổ chức này đơn giản cho quy mô đồ án, nhưng cần tách service nghiệp vụ nếu hệ thống tăng nhanh về độ phức tạp.")

    heading(doc, "Thiết kế dữ liệu", 2)
    p(doc, "Mã nguồn hiện tại định nghĩa 12 mô hình Mongoose. Các mô hình lõi User–Project–Task–Resource mô tả tài khoản, dự án, công việc và hồ sơ năng lực. Những mô hình còn lại bổ sung phiên đăng nhập, nhóm công việc, lịch lặp, kết quả tối ưu, thông báo, nhật ký và cấu hình tổ chức.")
    figure(doc, "data_model.png", "Hình 3-2 Mô hình dữ liệu khái quát")
    add_table(doc, ["Collection", "Dữ liệu chính", "Quan hệ tiêu biểu"], [
        ("Users", "Tài khoản, vai trò, quyền, quản lý trực tiếp", "Tham chiếu bởi Project, Task, Resource"),
        ("RefreshTokens", "Băm token, chuỗi phiên, thời điểm hết hạn", "Thuộc User"),
        ("Projects", "Mục tiêu, thời gian, thành viên, tiến độ", "Chứa nhiều Task và TaskGroup"),
        ("Tasks", "Effort, trạng thái, kỹ năng, phụ thuộc, cộng tác", "Thuộc Project; assignee là User"),
        ("TaskGroups", "Nhóm và thứ tự hiển thị công việc", "Thuộc Project"),
        ("RecurringTasks", "Quy tắc sinh công việc định kỳ", "Tạo Task theo lịch"),
        ("Resources", "Kỹ năng, capacity, FTE, lịch nghỉ, chi phí", "Liên kết User và Department"),
        ("Departments", "Phòng ban và thông tin quản lý", "Được Resource tham chiếu"),
        ("OptimizationResults", "Phương án, fitness, metrics, hội tụ", "Tham chiếu Task và Resource"),
        ("Notifications", "Nội dung, trạng thái đọc, liên kết điều hướng", "Thuộc người nhận User"),
        ("ActivityLogs", "Hành động, tác nhân, địa chỉ IP, snapshot", "Phục vụ truy vết"),
        ("CompanySettings", "Cấu hình tổ chức và ngày làm việc", "Một bản cấu hình dùng chung"),
    ], [3.3, 7.2, 5.5])
    caption(doc, "Bảng 3-2 Danh mục các collection", "table")

    heading(doc, "Thiết kế API và giao tiếp thời gian thực", 2)
    p(doc, "Ứng dụng tổ chức 12 nhóm route dưới tiền tố /api: auth, projects, tasks, resources, departments, optimization, analytics, notifications, activity-logs, task-groups, recurring-tasks và company-settings. Tại thời điểm lập báo cáo, mã nguồn chứa hơn 100 khai báo endpoint, bao phủ CRUD, truy vấn thống kê, nhập Excel, benchmark, phiên đăng nhập và các thao tác cộng tác.")
    add_table(doc, ["Phương thức/nhóm", "Ví dụ", "Mục đích"], [
        ("GET", "/api/projects, /api/tasks, /api/analytics/dashboard", "Đọc danh sách, chi tiết và thống kê"),
        ("POST", "/api/tasks, /api/optimization/run/hybrid", "Tạo dữ liệu hoặc khởi chạy xử lý"),
        ("PUT/PATCH", "/api/tasks/:id, /api/tasks/:id/status", "Cập nhật toàn phần hoặc một trạng thái"),
        ("DELETE", "/api/departments/:id, /api/tasks/:id", "Xóa có kiểm tra quyền và ràng buộc"),
        ("Socket.IO", "notification:new, notification:read", "Đẩy thay đổi theo thời gian thực"),
    ], [3.0, 7.0, 6.0])
    caption(doc, "Bảng 3-3 Thiết kế giao tiếp của hệ thống", "table")

    heading(doc, "Thiết kế bảo mật và phân quyền", 2)
    p(doc, "Mỗi request bảo vệ đi qua xác thực JWT. Sau xác thực, middleware kiểm tra vai trò, quyền ứng dụng hoặc quyền trên bản ghi cụ thể. Đối với công việc, middleware phân biệt quyền tạo, sửa, xóa, đổi trạng thái, báo cáo kết quả, nhân bản, di chuyển, cập nhật hạn và quản lý checklist/follower. Các thao tác quan trọng được ghi ActivityLog để phục vụ kiểm tra sau sự cố.")
    p(doc, "Luồng làm mới token dùng cơ chế single-flight ở client: khi nhiều request đồng thời nhận 401, chỉ một request refresh được gửi; các request còn lại chờ cùng kết quả. Nếu refresh thất bại, token trong bộ nhớ bị xóa và người dùng được đưa về màn hình đăng nhập. Ở server, refresh token được xoay vòng và phát hiện tái sử dụng để thu hồi cả chuỗi phiên.")

    heading(doc, "Thiết kế giao diện", 2)
    p(doc, "Giao diện được tổ chức quanh sidebar và header dùng chung; nội dung trang được nạp lười theo route để giữ khung ứng dụng ổn định trong khi tải chunk. Dashboard ưu tiên chỉ số hành động: dự án hoạt động, công việc trễ/bị chặn/chưa phân công, nhân sự quá tải và công việc cập nhật gần đây. Trang Công việc hỗ trợ bảng và Kanban; trang Lịch hỗ trợ ngày/tuần/tháng; Gantt biểu diễn phụ thuộc và đường găng; trang Tối ưu hóa hiển thị cấu hình, phương án, chỉ số và lịch sử.")
    p(doc, "Ngôn ngữ và định dạng được tách khỏi mã trạng thái. Hai tệp locale Việt/Anh có cùng tập khóa; nhãn enum được ánh xạ riêng để giá trị lưu trong cơ sở dữ liệu không phụ thuộc ngôn ngữ hiển thị. Theme sáng/tối dùng token thiết kế chung nhằm duy trì độ tương phản và tính nhất quán.")


def build_chapter_4(doc):
    heading(doc, ": HIỆN THỰC – KIỂM THỬ", 1)
    heading(doc, "Môi trường và cấu trúc hiện thực", 2)
    add_table(doc, ["Thành phần", "Công nghệ", "Phiên bản trong project"], [
        ("Web frontend", "React, Vite, Ant Design, Axios", "18.3; 5.4; 6.6; 1.7"),
        ("Backend", "Node.js, Express, Socket.IO", "Node ≥ 18; 4.19; 4.8"),
        ("Dữ liệu", "MongoDB, Mongoose", "7.x; 8.5"),
        ("Bảo mật", "JWT, bcryptjs, Helmet, rate-limit", "9.0; 2.4; 8.3; 8.6"),
        ("Mobile bổ trợ", "React Native, Expo", "0.81; 54"),
        ("Kiểm thử", "Node assertions, Vitest, Testing Library", "Vitest 2.1"),
    ], [4.0, 7.0, 5.0])
    caption(doc, "Bảng 4-1 Môi trường công nghệ", "table")
    p(doc, "Repository được tổ chức theo mô hình monorepo gồm client, server, mobile và docs. Frontend chia page, component, context, service, i18n và utility. Backend chia route, controller, model, middleware, service, algorithm, analytics, config và utility. Cấu trúc này làm rõ ranh giới giữa giao diện, truy cập API, nghiệp vụ server và thuật toán.")

    heading(doc, "Hiện thực các mô-đun nghiệp vụ", 2)
    heading(doc, "Xác thực, tài khoản và phiên đăng nhập", 3)
    p(doc, "Người dùng có thể đăng ký, đăng nhập, đổi mật khẩu, cập nhật hồ sơ, xem phiên đang hoạt động, thu hồi một phiên hoặc đăng xuất mọi thiết bị. Quản trị viên quản lý trạng thái, vai trò, email, quản lý trực tiếp, quyền ứng dụng, quyền đặc biệt và đặt lại mật khẩu. Access token có thời hạn ngắn; refresh token httpOnly giảm nguy cơ bị đọc bởi mã JavaScript độc hại.")
    heading(doc, "Dự án, công việc và cộng tác", 3)
    p(doc, "Mô-đun dự án hỗ trợ lọc, tìm kiếm, thành viên, allocation và chi tiết theo tab. Công việc hỗ trợ bảng/Kanban, Excel import, nhóm, công việc con, phụ thuộc chống chu trình, nhắc việc, nhân bản, di chuyển giữa nhóm/dự án, cập nhật hạn, bình luận, checklist, follower và báo cáo kết quả. Tiến độ dự án được tính lại khi dữ liệu công việc thay đổi.")
    heading(doc, "Nhân sự, lịch nghỉ và phân tích tải", 3)
    p(doc, "Hồ sơ Resource liên kết với User, chứa mã nhân viên tự sinh, phòng ban, chức danh, kỹ năng mức 1–4, năng lực tuần, FTE, đơn giá, trạng thái và các kỳ nghỉ. Người dùng tự quản lý lịch nghỉ của mình; quản lý có thể cập nhật dữ liệu trong phạm vi quyền. Workload trend trải số giờ theo ngày làm việc, xét FTE và lịch nghỉ rồi gộp theo tuần để phát hiện quá tải.")
    heading(doc, "Lịch, Gantt, báo cáo và thông báo", 3)
    p(doc, "Lịch hỗ trợ ba chế độ ngày, tuần và tháng. Gantt tính thời lượng, mốc, đồ thị phụ thuộc, chu trình và đường găng. Báo cáo tổng hợp utilization, burnout risk, phân bố trạng thái, số giờ và cho phép xuất dữ liệu. Khi giao việc hoặc thay đổi liên quan, server lưu Notification rồi phát sự kiện Socket.IO đến đúng phòng người dùng; giao hàng loạt được gộp thông báo theo người nhận để tránh spam.")
    heading(doc, "Công việc lặp và cấu hình công ty", 3)
    p(doc, "RecurringTask lưu mẫu công việc và quy tắc lặp, hỗ trợ xem trước lịch sinh, kích hoạt ngay và quản lý vòng đời. CompanySetting tập trung cấu hình tổ chức và ngày làm việc, tạo nền tảng để thống nhất cách tính lịch thay vì viết cứng ở nhiều màn hình.")

    heading(doc, "Hiện thực bộ máy tối ưu hóa", 2)
    heading(doc, "Genetic Algorithm", 3)
    p(doc, "Lớp GeneticAlgorithm nhận danh sách task/resource và cấu hình. Quần thể được sinh trong miền hợp lệ nếu có; mỗi vòng lặp đánh giá fitness, sắp xếp, giữ elite, chọn cha mẹ theo tournament, lai ghép uniform và đột biến. Lịch sử hội tụ được lấy mẫu để hiển thị; kết quả cuối bao gồm assignments, fitness, workload variance, utilization, skill match, overallocation, total cost và execution time.")
    heading(doc, "CSP Solver", 3)
    p(doc, "CSPSolver xây đồ thị xung đột từ các cặp công việc phụ thuộc và chồng lịch. Miền được lọc bởi skill/availability, sau đó node consistency loại lựa chọn không đủ năng lực cho riêng công việc. AC-3 lan truyền trên đồ thị ràng buộc; backtracking dùng MRV và LCV, cập nhật workload tạm thời và hoàn tác khi nhánh thất bại. Bộ giải giới hạn số vòng và thời gian để tránh khóa request quá lâu.")
    heading(doc, "Hybrid và vòng đời kết quả", 3)
    p(doc, "Ở chế độ Hybrid, miền khả thi do CSP tạo được truyền vào GA. Kết quả ghi số cặp task–resource trước/sau lọc và số task phải mở lại miền. OptimizationResult trải qua các trạng thái running, completed hoặc failed; phương án hoàn thành có thể so sánh, áp dụng và rollback. Khi áp dụng, hệ thống gán assignee theo resource.user, gửi thông báo và ghi nhật ký.")

    heading(doc, "Chiến lược kiểm thử", 2)
    p(doc, "Project sử dụng bốn lớp kiểm chứng: kiểm thử logic thuần, kiểm thử component, kiểm thử tích hợp API/Socket và build production. Backend test runner tạo database kiểm thử riêng, seed lại trước mỗi suite và chạy server trên cổng riêng; nhờ đó không làm thay đổi dữ liệu phát triển. Các suite bao phủ security, scoring, workload trend, CSP, Hybrid, API, chi tiết dự án và Socket.IO.")
    add_table(doc, ["Nhóm", "Phạm vi", "Kết quả lần chạy lập báo cáo"], [
        ("Logic Gantt", "Thời lượng, mốc, CPM, chu trình, lựa chọn predecessor", "33/33 đạt"),
        ("i18n và enum", "Khóa Việt/Anh, nội suy, enum khớp schema", "12/12 đạt"),
        ("Component web", "Refresh token, route, dashboard, resource, notification", "27/31 đạt; 4 ca chưa đạt"),
        ("Build production", "Biên dịch và chia chunk Vite", "Đạt; 4.218 module, 25,58 giây"),
        ("Backend tích hợp", "API, Socket, database test", "Chưa chạy đủ trong môi trường lập báo cáo"),
    ], [3.2, 8.0, 4.8])
    caption(doc, "Bảng 4-2 Kết quả kiểm chứng", "table")
    figure(doc, "verification.png", "Hình 4-1 Tổng hợp kết quả kiểm chứng")
    p(doc, "Bốn ca component chưa đạt không làm build thất bại nhưng phải được xem là tồn đọng chất lượng. Log cho thấy một nhóm kỳ vọng truy vấn accessible name không còn khớp cấu trúc giao diện Dashboard sau thay đổi; ngoài ra môi trường jsdom phát cảnh báo điều hướng/getComputedStyle và các API Ant Design đã deprecated. Trước nghiệm thu cuối cần cập nhật test theo hành vi người dùng, xử lý cảnh báo có ảnh hưởng và chạy lại toàn bộ suite.")

    heading(doc, "Thực nghiệm thuật toán", 2)
    p(doc, "Tài liệu thực nghiệm trong repository mô tả bộ dữ liệu 20 công việc × 12 nhân sự, mỗi người tập trung vào một kỹ năng. CSP giảm số cặp có thể gán từ 240 xuống 48, tương đương giảm 80% không gian miền. Kết quả trung bình 40 lần chạy cho thấy GA độc lập đạt fitness 0,8519 và dừng sau 119 thế hệ; Hybrid đạt fitness 0,8532 và dừng sau 90 thế hệ.")
    add_table(doc, ["Phương pháp", "Fitness trung bình", "Số thế hệ tới khi dừng", "Nhận xét"], [
        ("GA", "0,8519", "119", "Không bị giới hạn miền, cần nhiều thế hệ để loại lựa chọn kém"),
        ("Hybrid", "0,8532", "90", "Fitness tương đương, giảm khoảng 24% số thế hệ"),
    ], [3.0, 3.5, 4.0, 5.5])
    caption(doc, "Bảng 4-3 So sánh GA và Hybrid", "table")
    p(doc, "Kết quả cho thấy đóng góp chính của Hybrid là tốc độ hội tụ, không phải mức tăng fitness lớn. Đây là kết quả hợp lý: GA có thể tự học tránh nhân sự thiếu kỹ năng, nhưng phải trả giá bằng nhiều lần đánh giá; CSP loại các giá trị chắc chắn không phù hợp ngay trước khi tiến hóa.")

    heading(doc, "Đánh giá kết quả và rủi ro kỹ thuật", 2)
    add_table(doc, ["Nội dung", "Đánh giá"], [
        ("Độ bao phủ nghiệp vụ", "Bao phủ chu trình từ lập dự án, công việc, nhân sự đến tối ưu, báo cáo và truy vết."),
        ("Tính giải thích", "Phương án kèm skill match, tải, chi phí, quá tải và báo cáo ràng buộc."),
        ("Khả năng mở rộng", "Route-level lazy loading và module hóa tốt; chunk entry vẫn lớn và controller còn gọi model trực tiếp."),
        ("Bảo mật", "Có token rotation, httpOnly cookie, Helmet, CORS, rate limit, validation và activity log."),
        ("Chất lượng", "Build và logic cốt lõi đạt; component còn 4 ca chưa đạt, backend cần chạy đầy đủ với MongoDB test."),
        ("Thuật toán", "Giải gần đúng phù hợp quy mô đồ án; chưa chứng minh tối ưu toàn cục và phụ thuộc chất lượng dữ liệu đầu vào."),
    ], [4.0, 12.0])
    caption(doc, "Bảng 4-4 Đánh giá hệ thống", "table")


def build_chapter_5(doc):
    heading(doc, ": KẾT LUẬN", 1)
    heading(doc, "Kết quả đạt được", 2)
    p(doc, "Đồ án đã xây dựng được hệ thống quản lý luồng công việc đa dự án có khả năng tập trung dữ liệu dự án, công việc, nhân sự, kỹ năng, lịch nghỉ và mức sử dụng nguồn lực. Hệ thống cung cấp nhiều góc nhìn gồm dashboard, bảng, Kanban, lịch, Gantt và báo cáo; hỗ trợ phân quyền, thông báo thời gian thực, nhật ký hoạt động, đa ngôn ngữ và giao diện sáng/tối.")
    p(doc, "Đóng góp nổi bật là tích hợp ba phương pháp phân bổ nhân sự trong cùng quy trình nghiệp vụ. GA tối ưu đa mục tiêu, CSP xử lý ràng buộc cứng và Hybrid dùng CSP thu hẹp miền trước khi tiến hóa. Phương án không chỉ đưa ra người được giao mà còn kèm các chỉ số, lịch sử, so sánh, cảnh báo nới ràng buộc, thao tác áp dụng và hoàn tác. Thực nghiệm cho thấy Hybrid giảm khoảng 24% số thế hệ trong bộ dữ liệu mô phỏng mà vẫn giữ chất lượng fitness tương đương.")
    p(doc, "Về kỹ thuật phần mềm, project có cấu trúc client/server/mobile rõ ràng, API theo REST, mô hình dữ liệu Mongoose, refresh token có thể thu hồi, kiểm tra quyền theo vai trò và bản ghi, bộ test nhiều lớp và quy trình build production. Báo cáo đã được cập nhật dựa trực tiếp trên trạng thái mã nguồn hiện tại, khắc phục nhiều mô tả cũ không còn đúng.")

    heading(doc, "Hạn chế của đồ án", 2)
    bullet(doc, "Kết quả tối ưu là nghiệm gần đúng; GA không bảo đảm tối ưu toàn cục và nhạy với tham số/quy mô dữ liệu.")
    bullet(doc, "Ràng buộc kỹ năng CSP dùng điểm tổng hợp ≥ 0,5, nên có thể chấp nhận nhân sự thiếu một kỹ năng nếu các kỹ năng khác bù điểm.")
    bullet(doc, "AC-3 trên ràng buộc ≠ không phát hiện được mọi trường hợp vô nghiệm kiểu all-different; lịch của công việc là dữ liệu cố định, thuật toán chưa đồng thời tối ưu thời gian.")
    bullet(doc, "Chất lượng phương án phụ thuộc dữ liệu kỹ năng, effort, capacity và lịch nghỉ do người dùng nhập.")
    bullet(doc, "Bộ kiểm thử component còn 4 ca chưa đạt tại thời điểm lập báo cáo; bộ tích hợp backend cần được chạy lại đầy đủ trong môi trường MongoDB test trước nghiệm thu.")
    bullet(doc, "Frontend còn chunk đầu vào lớn; backend chưa tách tầng service nghiệp vụ hoàn chỉnh; triển khai quy mô lớn cần hàng đợi xử lý thuật toán và quan sát hệ thống tốt hơn.")

    heading(doc, "Hướng phát triển", 2)
    bullet(doc, "Mở rộng mô hình thành bài toán đồng tối ưu người thực hiện, lịch bắt đầu/kết thúc và phụ thuộc; nghiên cứu NSGA-II hoặc tối ưu Pareto để không ép mọi mục tiêu vào một trọng số cố định.")
    bullet(doc, "Bổ sung ràng buộc all-different toàn cục, kỹ năng bắt buộc theo từng mục và cơ chế giải thích nguyên nhân một nhân sự bị loại khỏi miền.")
    bullet(doc, "Xây dựng hàng đợi tác vụ nền cho benchmark/optimization, cache kết quả và giám sát thời gian, bộ nhớ, lỗi theo lần chạy.")
    bullet(doc, "Hoàn thiện kiểm thử end-to-end trên trình duyệt, kiểm thử tải, kiểm thử bảo mật và pipeline CI/CD; đặt ngưỡng chất lượng bắt buộc trước khi hợp nhất mã nguồn.")
    bullet(doc, "Giảm bundle bằng tách nhỏ thành phần Ant Design, đo Core Web Vitals và tối ưu trải nghiệm mobile; đồng bộ thêm nghiệp vụ vào ứng dụng React Native.")
    bullet(doc, "Tích hợp lịch/nhân sự của bên thứ ba, tệp đính kèm, @mention, audit bất biến và báo cáo quản trị nâng cao khi triển khai thực tế.")


def build_references_appendix(doc):
    heading(doc, "TÀI LIỆU THAM KHẢO", 1)
    refs = [
        "[1] J. Błażewicz, J. K. Lenstra và A. H. G. Rinnooy Kan. Scheduling subject to resource constraints: classification and complexity. Discrete Applied Mathematics, 1983.",
        "[2] R. Kolisch và S. Hartmann. Experimental investigation of heuristics for resource-constrained project scheduling. European Journal of Operational Research, 2006.",
        "[3] H. W. Kuhn. The Hungarian method for the assignment problem. Naval Research Logistics Quarterly, 1955.",
        "[4] J. H. Holland. Adaptation in Natural and Artificial Systems. University of Michigan Press, 1975.",
        "[5] K. Deb. Multi-Objective Optimization Using Evolutionary Algorithms. Wiley, 2001.",
        "[6] S. Russell và P. Norvig. Artificial Intelligence: A Modern Approach, 4th Edition. Pearson, 2020.",
        "[7] Project Management Institute. A Guide to the Project Management Body of Knowledge (PMBOK Guide), 7th Edition. PMI, 2021.",
        "[8] Atlassian. Jira Software Cloud Documentation. https://support.atlassian.com/jira-software-cloud/ (truy cập tháng 8/2026).",
        "[9] Microsoft. Microsoft Project Documentation. https://support.microsoft.com/project (truy cập tháng 8/2026).",
        "[10] Asana. Workload and Resource Management. https://asana.com/resources/resource-management (truy cập tháng 8/2026).",
        "[11] Meta Open Source. React Documentation. https://react.dev/ (truy cập tháng 8/2026).",
        "[12] MongoDB Inc. MongoDB Manual. https://www.mongodb.com/docs/manual/ (truy cập tháng 8/2026).",
        "[13] OpenJS Foundation. Node.js Documentation. https://nodejs.org/docs/ (truy cập tháng 8/2026).",
        "[14] OWASP Foundation. OWASP Application Security Verification Standard. https://owasp.org/www-project-application-security-verification-standard/ (truy cập tháng 8/2026).",
    ]
    for ref in refs:
        para = doc.add_paragraph(style="Tai lieu tham khao" if "Tai lieu tham khao" in doc.styles else "Normal")
        para.paragraph_format.first_line_indent = Cm(-0.8)
        para.paragraph_format.left_indent = Cm(0.8)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.add_run(ref)

    heading(doc, "PHỤ LỤC", 1)
    heading(doc, "A. Hướng dẫn cài đặt và chạy hệ thống", 2)
    p(doc, "Yêu cầu Node.js từ phiên bản 18, MongoDB đang hoạt động và các biến môi trường phù hợp. Quy trình tham khảo:")
    steps = [
        "Cài dependency cho toàn bộ workspace bằng lệnh npm run install:all tại thư mục gốc.",
        "Sao chép .env.example thành .env; thiết lập MONGODB_URI, JWT_SECRET, CLIENT_URL và các biến email nếu sử dụng.",
        "Tạo dữ liệu mẫu bằng npm run seed --prefix server. Lưu ý thao tác seed xóa dữ liệu của các collection mục tiêu; chỉ dùng trên cơ sở dữ liệu phát triển.",
        "Chạy đồng thời web và server bằng npm run dev; web mặc định ở cổng 5173, API ở cổng 5000.",
        "Kiểm tra build production bằng npm run build và chạy kiểm thử ở từng package trước khi triển khai.",
    ]
    for item in steps: numbered(doc, item)

    heading(doc, "B. Danh mục biến môi trường chính", 2)
    add_table(doc, ["Biến", "Mục đích", "Ghi chú"], [
        ("PORT", "Cổng HTTP backend", "Mặc định 5000"),
        ("NODE_ENV", "Môi trường chạy", "production bắt buộc có JWT_SECRET"),
        ("MONGODB_URI", "Chuỗi kết nối MongoDB", "Nên tách database dev/test/prod"),
        ("JWT_SECRET", "Khóa ký token", "Không lưu khóa thật trong repository"),
        ("CLIENT_URL", "Danh sách origin được CORS chấp nhận", "Có thể ngăn cách bằng dấu phẩy"),
        ("AUTH_RATE_LIMIT_MAX", "Giới hạn đăng nhập sai", "Mặc định 10/15 phút/IP"),
        ("API_RATE_LIMIT_MAX", "Giới hạn request API", "Mặc định 1000/15 phút/IP"),
        ("VITE_API_URL", "Base URL từ web đến API", "Có thể dùng /api qua Vite proxy"),
    ], [4.0, 7.0, 5.0])
    caption(doc, "Bảng PL-1 Các biến môi trường chính", "table")

    heading(doc, "C. Danh mục kiểm thử cần chạy trước nghiệm thu", 2)
    add_table(doc, ["Bước", "Lệnh/hoạt động", "Tiêu chí"], [
        ("1", "Build web production", "Không có lỗi biên dịch"),
        ("2", "Kiểm thử logic và component client", "Tất cả ca đạt; không bỏ qua ca thất bại"),
        ("3", "Khởi động MongoDB test và chạy server test", "Security, API, project detail, socket đều đạt"),
        ("4", "Chạy riêng CSP, Hybrid, scoring, workload trend", "Kết quả xác định, không hồi quy"),
        ("5", "Kiểm thử thủ công 3 vai trò", "Đúng ranh giới quyền và luồng áp dụng/rollback"),
        ("6", "Rà soát tài liệu, mục lục, caption", "Ctrl+A, F9; không còn Error! Bookmark"),
    ], [2.0, 8.0, 6.0])
    caption(doc, "Bảng PL-2 Danh sách kiểm tra trước nghiệm thu", "table")

    heading(doc, "D. Ghi chú cập nhật mục lục trong Microsoft Word", 2)
    p(doc, "Tài liệu chứa các trường tự động cho Mục lục, Danh mục hình ảnh và Danh mục bảng biểu. Khi mở file bằng Microsoft Word, chọn toàn bộ tài liệu bằng Ctrl+A, nhấn F9 và chọn Update entire table. Sau khi cập nhật, kiểm tra lại số trang, ngắt trang và mọi tham chiếu trước khi in hoặc nộp.")


def finalize(doc):
    # Ensure fields update when Word opens.
    settings = doc.settings._element
    update = settings.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields')
        settings.append(update)
    update.set(qn('w:val'), 'true')

    # Apply font consistently to newly created and existing body text without disturbing cover sizes.
    for para in doc.paragraphs[20:]:
        for run in para.runs:
            run.font.name = "Times New Roman"
            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    doc.core_properties.title = "Xây dựng phần mềm quản lý luồng công việc đa dự án và tối ưu hóa phân bổ nhân sự"
    doc.core_properties.subject = "Tài liệu đồ án tốt nghiệp – Resource Allocation Optimization"
    doc.core_properties.author = "Trần Xuân Trường; Lê Minh Tiến"
    doc.core_properties.keywords = "Resource Allocation, Multi-Project, Genetic Algorithm, CSP, Hybrid"
    doc.save(OUTPUT)


def main():
    make_diagrams()
    doc = Document(SOURCE)
    normalize_styles(doc)
    update_existing(doc)
    insert_frontmatter(doc)
    build_chapter_2(doc)
    build_chapter_3(doc)
    build_chapter_4(doc)
    build_chapter_5(doc)
    build_references_appendix(doc)
    finalize(doc)
    print(OUTPUT)


if __name__ == "__main__":
    main()
