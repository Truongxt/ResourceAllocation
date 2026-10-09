"""Add sample-style use-case specifications and UML diagrams to the thesis.

Nhom03.pdf is a presentation reference only.  All report facts and figures here
come from the ResourceAllocation project and its editable .vpp UML file.
"""

from __future__ import annotations

import re
import sqlite3
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(r"D:\ResourceAllocation\docs")
SOURCE = ROOT / "218_110_Truong_Tien_TaiLieu_HoanThien.docx"
OUTPUT = ROOT / "218_110_Truong_Tien_TaiLieu_BoSungUML_HoanChinh.docx"
VPP = ROOT / "ResourceAllocation_UML_MotKhung.vpp"
IMAGES = ROOT / "report_uml_images_ascii"

SPECS = [
    {
        "code": "UC-TH01", "title": "Đăng nhập và làm mới phiên", "links": "UC-02.02, UC-02.03",
        "actor": "Người dùng", "aux": "Hệ thống xác thực",
        "summary": "Xác thực tài khoản, cấp access token và duy trì phiên bằng refresh token xoay vòng.",
        "pre": "Tài khoản tồn tại và còn hoạt động; người dùng ở màn hình đăng nhập.",
        "post": "Phiên hợp lệ được tạo; client nhận access token, refresh token được quản lý theo loại client.",
        "flow": [
            ("Nhập email và mật khẩu.", "Kiểm tra định dạng dữ liệu đầu vào."),
            ("Gửi yêu cầu đăng nhập.", "Tra cứu tài khoản, so khớp mật khẩu đã băm và kiểm tra trạng thái."),
            ("Nhận kết quả đăng nhập.", "Cấp access token, lưu bản băm refresh token và trả hồ sơ người dùng."),
            ("Tiếp tục thao tác khi access token hết hạn.", "Kiểm tra refresh token, xoay vòng token và cấp access token mới."),
        ],
        "alt": "Nếu email hoặc mật khẩu sai, hệ thống từ chối đăng nhập và cho phép nhập lại.",
        "exc": "Refresh token hết hạn, bị thu hồi hoặc bị tái sử dụng: hủy phiên và yêu cầu đăng nhập lại.",
        "source": "auth.routes.js, auth.controller.js; POST /api/auth/login, POST /api/auth/refresh",
        "act": "ACT01", "seq": "SEQ01",
    },
    {
        "code": "UC-TH02", "title": "Tạo dự án và quản lý thành viên", "links": "UC-03.02, UC-03.03",
        "actor": "Quản lý dự án / Quản trị viên", "aux": "Thành viên dự án",
        "summary": "Khởi tạo dự án, thiết lập thông tin quản lý và bổ sung thành viên với tỷ lệ tham gia.",
        "pre": "Người thao tác đã đăng nhập và có quyền tạo hoặc quản lý dự án.",
        "post": "Dự án và danh sách thành viên được lưu; thao tác quản trị được ghi nhận.",
        "flow": [
            ("Nhập tên, thời gian, độ ưu tiên và dữ liệu dự án.", "Kiểm tra dữ liệu và quyền tạo dự án."),
            ("Xác nhận tạo dự án.", "Lưu Project, gắn người tạo/quản lý và trả thông tin dự án mới."),
            ("Chọn người tham gia, vai trò và tỷ lệ phân bổ.", "Kiểm tra tài khoản, cập nhật members của dự án."),
            ("Mở lại dự án để theo dõi.", "Hiển thị thành viên và tiến độ theo phạm vi truy cập."),
        ],
        "alt": "Có thể tạo dự án trước và thêm thành viên sau; người không thuộc dự án chỉ thấy dữ liệu theo quyền của mình.",
        "exc": "Dữ liệu không hợp lệ hoặc thiếu quyền: từ chối lưu, giữ nguyên dự án/thành viên hiện tại.",
        "source": "project.routes.js, project.controller.js; /api/projects, /api/projects/:id/members",
        "act": "ACT02", "seq": "SEQ06",
    },
    {
        "code": "UC-TH03", "title": "Tạo, cập nhật công việc và cộng tác", "links": "UC-03.06–UC-03.13",
        "actor": "Quản lý dự án / Thành viên được cấp quyền", "aux": "Người được giao việc",
        "summary": "Tạo công việc, thiết lập ràng buộc, gán người thực hiện và cập nhật trạng thái trên Kanban.",
        "pre": "Dự án tồn tại; người thao tác có quyền trên dự án/công việc liên quan.",
        "post": "Task được lưu; trạng thái, tiến độ dự án và thông báo liên quan được cập nhật.",
        "flow": [
            ("Nhập tiêu đề, effort, thời hạn, kỹ năng và phụ thuộc.", "Kiểm tra quyền, dữ liệu và quan hệ phụ thuộc."),
            ("Xác nhận tạo/gán công việc.", "Lưu Task, cập nhật dự án và phát sự kiện task:created."),
            ("Kéo thả thẻ Kanban hoặc cập nhật tiến độ.", "Kiểm tra chuyển trạng thái, lưu thay đổi và phát task:updated."),
            ("Bình luận, quản lý checklist hoặc báo cáo kết quả.", "Lưu nội dung cộng tác và tính lại tiến độ/workload khi cần."),
        ],
        "alt": "Nếu thao tác Kanban thất bại, giao diện hoàn nguyên vị trí thẻ và hiển thị lỗi.",
        "exc": "Phụ thuộc tự tham chiếu, khác dự án hoặc tạo chu trình: từ chối lưu quan hệ phụ thuộc.",
        "source": "task.routes.js, task.controller.js; POST /api/tasks, PATCH /api/tasks/:id/status",
        "act": "ACT03", "seq": "SEQ02",
    },
    {
        "code": "UC-TH04", "title": "Quản lý hồ sơ nhân sự và lịch nghỉ", "links": "UC-04.02–UC-04.07",
        "actor": "Quản trị viên / Quản lý được cấp quyền", "aux": "Nhân sự",
        "summary": "Cập nhật năng lực, capacity/FTE và khoảng vắng mặt làm đầu vào cho phân bổ nguồn lực.",
        "pre": "Hồ sơ nhân sự tồn tại hoặc được tạo mới; người thao tác có quyền quản lý tương ứng.",
        "post": "Hồ sơ, kỹ năng và lịch nghỉ được lưu; workload khả dụng được tính lại.",
        "flow": [
            ("Tạo hoặc chọn hồ sơ nhân sự.", "Kiểm tra tài khoản liên kết và quyền truy cập."),
            ("Nhập phòng ban, chức danh, capacity/FTE và ma trận kỹ năng.", "Kiểm tra dữ liệu năng lực, sau đó lưu hồ sơ."),
            ("Đăng ký hoặc quản lý kỳ nghỉ.", "Kiểm tra khoảng ngày và cập nhật thời gian không khả dụng."),
            ("Xem tải phân công.", "Tính lại workload; cảnh báo khi tải vượt khả năng."),
        ],
        "alt": "Nhân sự có thể đăng ký lịch nghỉ cá nhân nếu được phép; quản trị viên quản lý lịch nghỉ trong phạm vi tổ chức.",
        "exc": "Khoảng ngày không hợp lệ hoặc dữ liệu capacity sai: không lưu và yêu cầu chỉnh sửa.",
        "source": "resource.routes.js, resource.controller.js; /api/resources",
        "act": "ACT04", "seq": None,
    },
    {
        "code": "UC-TH05", "title": "Chạy tối ưu hóa Hybrid CSP–GA", "links": "UC-05.01, UC-05.04, UC-05.05",
        "actor": "Quản lý dự án / Quản trị viên", "aux": "Bộ giải CSP, Genetic Algorithm",
        "summary": "Sinh phương án phân bổ khả thi rồi cải thiện chất lượng theo mục tiêu tối ưu.",
        "pre": "Có công việc mở, nhân sự hoạt động và dữ liệu kỹ năng/capacity đủ để kiểm tra readiness.",
        "post": "OptimizationResult ở trạng thái hoàn tất hoặc thất bại, kèm chỉ số và lý do tương ứng.",
        "flow": [
            ("Chọn dự án/phạm vi và thuật toán Hybrid.", "Lấy task và resource trong phạm vi được phép."),
            ("Gửi yêu cầu chạy tối ưu.", "Kiểm tra readiness và tạo bản ghi kết quả đang chạy."),
            ("Theo dõi quá trình giải.", "CSP tạo miền/phương án khả thi; GA tiến hóa và chấm điểm."),
            ("Xem phương án.", "Lưu assignments, metrics, constraint report và trả kết quả."),
        ],
        "alt": "Người dùng có thể chọn chạy GA hoặc CSP riêng thay vì Hybrid để so sánh kết quả.",
        "exc": "Thiếu dữ liệu hoặc không tìm được lời giải: lưu trạng thái thất bại và thông báo nguyên nhân.",
        "source": "optimization.routes.js, optimization.controller.js; /api/optimization/run/hybrid",
        "act": "ACT05", "seq": "SEQ04",
    },
    {
        "code": "UC-TH06", "title": "Áp dụng và hoàn tác phương án", "links": "UC-05.09, UC-05.10",
        "actor": "Quản lý dự án / Quản trị viên", "aux": "Người được phân công",
        "summary": "Chuyển assignments từ kết quả tối ưu sang công việc thực tế và có thể khôi phục trạng thái trước áp dụng.",
        "pre": "Kết quả tối ưu đã hoàn tất; người thao tác có quyền áp dụng trên phạm vi tương ứng.",
        "post": "Task được phân công theo phương án hoặc được khôi phục từ snapshot khi hoàn tác.",
        "flow": [
            ("Chọn kết quả và xác nhận áp dụng.", "Kiểm tra trạng thái kết quả và quyền áp dụng."),
            ("Yêu cầu cập nhật phân công.", "Lưu snapshot assignee cũ, gán resource.user cho các Task."),
            ("Theo dõi thông báo hoàn tất.", "Tính lại workload/tiến độ, ghi log và gửi thông báo."),
            ("Nếu cần, chọn hoàn tác.", "Khôi phục snapshot và cập nhật trạng thái kết quả."),
        ],
        "alt": "Người dùng có thể chỉ xem hoặc so sánh kết quả mà chưa áp dụng.",
        "exc": "Kết quả chưa hoàn tất, đã áp dụng hoặc thiếu quyền: từ chối thao tác; dữ liệu Task không bị đổi.",
        "source": "optimization.routes.js, optimization.controller.js; /api/optimization/:id/apply, /:id/rollback",
        "act": "ACT06", "seq": "SEQ05",
    },
    {
        "code": "UC-TH07", "title": "Quản lý và sinh công việc lặp", "links": "UC-06.11",
        "actor": "Quản lý dự án / Quản trị viên", "aux": "Bộ lập lịch",
        "summary": "Định nghĩa mẫu công việc và quy tắc lặp để hệ thống tạo occurrence theo lịch.",
        "pre": "Dự án còn hiệu lực; người thao tác có quyền tạo công việc trong dự án.",
        "post": "Mẫu lặp được lưu; đến hạn, task mới được tạo và mốc chạy tiếp theo được cập nhật.",
        "flow": [
            ("Nhập mẫu task và quy tắc lặp.", "Kiểm tra lịch, dự án và dữ liệu mẫu."),
            ("Xem trước rồi xác nhận lưu.", "Lưu RecurringTask và thời điểm nextRunAt."),
            ("Đến chu kỳ hoặc kích hoạt Run now.", "Kiểm tra lần sinh trước, tạo Task từ mẫu và cập nhật lastRun/nextRun."),
            ("Xem kết quả chu kỳ.", "Hiển thị task đã sinh và trạng thái xử lý."),
        ],
        "alt": "Mẫu có thể được tạm dừng; chu kỳ bị tạm dừng sẽ không sinh task mới.",
        "exc": "Quy tắc lịch không hợp lệ hoặc dự án không còn khả dụng: bỏ qua lần sinh và ghi nhận lý do.",
        "source": "recurringTask.routes.js, recurringTask.controller.js; /api/recurring-tasks",
        "act": "ACT07", "seq": "SEQ08",
    },
    {
        "code": "UC-TH08", "title": "Nhận và đánh dấu thông báo", "links": "UC-06.07, UC-06.08",
        "actor": "Thành viên", "aux": "Socket.IO, người tạo sự kiện",
        "summary": "Đưa sự kiện liên quan công việc đến đúng người nhận và đồng bộ trạng thái đã đọc.",
        "pre": "Người nhận có tài khoản; sự kiện nghiệp vụ được phát sinh.",
        "post": "Notification được lưu; client đang kết nối nhận sự kiện và trạng thái đã đọc được cập nhật.",
        "flow": [
            ("Đăng nhập và kết nối kênh thời gian thực.", "Xác thực token rồi cho client tham gia room user:id."),
            ("Có sự kiện gán/cập nhật công việc.", "Lưu Notification và emit notification:new tới người nhận."),
            ("Mở danh sách thông báo.", "Hiển thị badge/toast và nội dung trong phạm vi người dùng."),
            ("Chọn đánh dấu đã đọc.", "Cập nhật isRead và đồng bộ trạng thái."),
        ],
        "alt": "Nếu client không trực tuyến, thông báo vẫn nằm trong cơ sở dữ liệu để xem sau.",
        "exc": "Kết nối không có JWT hợp lệ: không tham gia room; API thông báo vẫn yêu cầu xác thực.",
        "source": "notification.routes.js, notification.controller.js; /api/notifications",
        "act": "ACT08", "seq": "SEQ03",
    },
    {
        "code": "UC-TH09", "title": "Nhập công việc từ Excel", "links": "UC-03.15",
        "actor": "Quản lý dự án / Quản trị viên", "aux": "Hệ thống nhập dữ liệu",
        "summary": "Chuyển các dòng trong workbook thành công việc thuộc dự án, kèm kết quả và lỗi theo dòng.",
        "pre": "Người thao tác có quyền tạo task; file Excel đúng định dạng và dự án mục tiêu tồn tại.",
        "post": "Các dòng hợp lệ được nhập; hệ thống trả thống kê nhập và lỗi theo dòng.",
        "flow": [
            ("Tải mẫu và chọn workbook.", "Nhận file, đọc sheet và chuẩn hóa các cột."),
            ("Gửi yêu cầu nhập.", "Kiểm tra dự án, quyền và từng dòng dữ liệu."),
            ("Xem trước lỗi, sửa file nếu cần.", "Trả chi tiết dòng không hợp lệ để người dùng xử lý."),
            ("Xác nhận nhập các dòng hợp lệ.", "Tạo Task, tính lại dữ liệu liên quan và trả thống kê."),
        ],
        "alt": "Nếu có dòng sai, người dùng chỉnh workbook và gửi lại; các lỗi được gắn đúng số dòng.",
        "exc": "File không đọc được hoặc không có dòng hợp lệ: không tạo task, trả thông báo lỗi.",
        "source": "task.routes.js, task.controller.js; /api/tasks/excel/import",
        "act": "ACT09", "seq": "SEQ07",
    },
    {
        "code": "UC-TH10", "title": "Xem dashboard, báo cáo và xuất dữ liệu", "links": "UC-06.04–UC-06.06",
        "actor": "Quản lý dự án / Quản trị viên", "aux": "Thành viên (dữ liệu cá nhân)",
        "summary": "Tổng hợp chỉ số dự án, công việc, utilization và xu hướng workload theo quyền truy cập.",
        "pre": "Người dùng đăng nhập; dữ liệu dự án/công việc đã có trong phạm vi nhìn thấy.",
        "post": "Dashboard/báo cáo hiển thị số liệu theo bộ lọc; dữ liệu có thể xuất theo chức năng hỗ trợ.",
        "flow": [
            ("Chọn khoảng thời gian và bộ lọc.", "Xác định phạm vi dữ liệu theo vai trò và quyền dự án."),
            ("Mở dashboard hoặc báo cáo.", "Tổng hợp KPI task, dự án, utilization và xu hướng tải."),
            ("Xem biểu đồ và ngoại lệ.", "Trả số liệu/biểu đồ chỉ trong phạm vi được phép."),
            ("Chọn xuất dữ liệu.", "Sinh CSV hoặc mở chế độ in/lưu PDF tương ứng."),
        ],
        "alt": "Nếu bộ lọc không có dữ liệu, giao diện hiển thị trạng thái rỗng thay vì số liệu của phạm vi khác.",
        "exc": "API lỗi hoặc phiên hết hạn: hiển thị lỗi, không xuất dữ liệu chưa xác thực.",
        "source": "analytics.routes.js, analytics.controller.js; /api/analytics/dashboard",
        "act": "ACT10", "seq": "SEQ09",
    },
    {
        "code": "UC-TH11", "title": "Xem lịch, Gantt và đường găng", "links": "UC-06.01, UC-06.02",
        "actor": "Thành viên / Quản lý dự án", "aux": "Bộ tính đường găng CPM",
        "summary": "Trực quan hóa task theo ngày, tuần, tháng và quan hệ phụ thuộc trên Gantt.",
        "pre": "Người dùng có quyền xem dự án/công việc trong khoảng ngày đã chọn.",
        "post": "Lịch và Gantt hiển thị task, phụ thuộc, critical path và slack trong phạm vi truy cập.",
        "flow": [
            ("Chọn dự án và khoảng ngày.", "Truy vấn task và mốc thời gian theo phạm vi quyền."),
            ("Mở chế độ Calendar hoặc Gantt.", "Nạp quan hệ phụ thuộc và các cạnh liên quan."),
            ("Yêu cầu phân tích đường găng.", "Tính critical path/slack từ mạng phụ thuộc."),
            ("Xem kết quả trực quan.", "Hiển thị thanh công việc, phụ thuộc và điểm cần chú ý."),
        ],
        "alt": "Có thể chuyển giữa góc nhìn ngày, tuần, tháng mà không thay đổi dữ liệu gốc.",
        "exc": "Không có task trong bộ lọc: hiển thị lịch rỗng; quan hệ phụ thuộc không hợp lệ cần được xử lý trước khi phân tích.",
        "source": "task.routes.js, task.controller.js; /api/tasks theo bộ lọc dự án/khoảng ngày",
        "act": None, "seq": "SEQ10",
    },
]

ACTIVITY_TITLES = {
    "ACT01": "Đăng nhập và làm mới phiên", "ACT02": "Vòng đời dự án",
    "ACT03": "Vòng đời công việc và cộng tác", "ACT04": "Quản lý nhân sự và lịch nghỉ",
    "ACT05": "Chạy tối ưu hóa", "ACT06": "Áp dụng và hoàn tác phương án",
    "ACT07": "Công việc lặp", "ACT08": "Thông báo thời gian thực",
    "ACT09": "Nhập công việc từ Excel", "ACT10": "Báo cáo và xuất dữ liệu",
}
SEQUENCE_TITLES = {
    "SEQ01": "Đăng nhập và refresh token", "SEQ02": "Tạo và cập nhật công việc",
    "SEQ03": "Thông báo thời gian thực", "SEQ04": "Chạy Hybrid CSP–GA",
    "SEQ05": "Áp dụng và hoàn tác tối ưu", "SEQ06": "Quản lý dự án và thành viên",
    "SEQ07": "Import Excel công việc", "SEQ08": "Sinh công việc lặp",
    "SEQ09": "Tải dashboard và báo cáo", "SEQ10": "Lịch và Gantt",
}
CLASS_TITLES = {
    "CLS01": "Mô hình miền và dữ liệu", "CLS02": "Kiến trúc lớp backend",
    "CLS03": "Kiến trúc lớp frontend",
}


def set_text(cell, value: str, bold: bool = False, size: float = 9.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(value))
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade(cell, color: str = "D9EAF7") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), color)


def no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def paragraph(doc, text: str = "", style: str = "Normal"):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    if style == "Normal":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def caption(doc, text: str, kind: str):
    p = paragraph(doc, text, "Figure Caption" if kind == "figure" else "Table Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_figure(doc, code: str, label: str, counter: list[int], page_break: bool = True) -> None:
    path = IMAGES / f"{code}.png"
    if not path.exists():
        raise FileNotFoundError(path)
    with Image.open(path) as image:
        ratio = image.width / image.height
    max_width_cm, max_height_cm = 16.0, 18.5
    width_cm = min(max_width_cm, max_height_cm * ratio)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.page_break_before = page_break
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    caption(doc, f"Hình 3-{counter[0]} {label}", "figure")
    counter[0] += 1


def diagram_catalog() -> list[tuple[str, str, str, str]]:
    conn = sqlite3.connect(VPP)
    conn.row_factory = sqlite3.Row
    output = []
    for diagram in conn.execute(
        "SELECT * FROM DIAGRAM WHERE NAME LIKE 'UC-0%' AND NAME NOT LIKE 'UC-00%' ORDER BY NAME"
    ).fetchall():
        did = diagram["ID"]
        rows = conn.execute(
            "SELECT * FROM DIAGRAM_ELEMENT WHERE DIAGRAM_ID=? ORDER BY rowid", (did,)
        ).fetchall()
        actors = {
            row["ID"]: conn.execute(
                "SELECT NAME FROM MODEL_ELEMENT WHERE ID=?", (row["MODEL_ELEMENT_ID"],)
            ).fetchone()[0]
            for row in rows if row["SHAPE_TYPE"] == "Actor"
        }
        case_actor = {}
        for row in rows:
            if row["SHAPE_TYPE"] != "Association":
                continue
            definition = row["DEFINITION"].decode("utf-8")
            actor = re.search(r"_fromShape=<[^:>]+:([^>]+)>", definition)
            case = re.search(r"_toShape=<[^:>]+:([^>]+)>", definition)
            if actor and case:
                case_actor[case.group(1)] = actors.get(actor.group(1), "Theo phân quyền")
        cases = [row for row in rows if row["SHAPE_TYPE"] == "UseCase"]
        for index, row in enumerate(cases, start=1):
            name = conn.execute(
                "SELECT NAME FROM MODEL_ELEMENT WHERE ID=?", (row["MODEL_ELEMENT_ID"],)
            ).fetchone()[0]
            code = f"{diagram['NAME'][:5]}.{index:02d}"
            output.append((code, name, case_actor.get(row["ID"], "Theo phân quyền"), diagram["NAME"][6:]))
    conn.close()
    if len(output) != 72:
        raise ValueError(f"Expected 72 use cases, found {len(output)}")
    return output


def add_catalog(doc) -> None:
    rows = diagram_catalog()
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ("Mã", "Use case", "Tác nhân chính", "Phân hệ")
    widths = (Cm(2.6), Cm(5.0), Cm(3.5), Cm(4.8))
    for index, width in enumerate(widths):
        table.columns[index].width = width
    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        set_text(cell, heading, bold=True, size=9)
        shade(cell)
    repeat_header(table.rows[0])
    for code, title, actor, group in rows:
        cells = table.add_row().cells
        for index, value in enumerate((code.replace("-", "\u2011"), title, actor, group)):
            cells[index].width = widths[index]
            set_text(cells[index], value, size=8.5)
        no_wrap = OxmlElement("w:noWrap")
        cells[0]._tc.get_or_add_tcPr().append(no_wrap)
        no_split(table.rows[-1])
    caption(doc, "Bảng 3-4 Danh mục 72 use case và tác nhân liên quan", "table")


def add_spec_table(doc, spec: dict, table_number: int) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    def row(left: str, right: str, band: bool = False):
        cells = table.add_row().cells
        cells[0].width, cells[1].width = Cm(5.0), Cm(10.9)
        set_text(cells[0], left, bold=True)
        set_text(cells[1], right)
        if band:
            shade(cells[0]); shade(cells[1])
        no_split(table.rows[-1])

    row("Mã use case", f"{spec['code']} (liên quan {spec['links']})", True)
    row("Tên use case", spec["title"])
    row("Mô tả sơ lược", spec["summary"])
    row("Actor chính", spec["actor"])
    row("Actor phụ", spec["aux"])
    row("Tiền điều kiện", spec["pre"])
    row("Hậu điều kiện", spec["post"])
    cells = table.add_row().cells
    cells[0].merge(cells[1])
    set_text(table.rows[-1].cells[0], "Luồng sự kiện chính", bold=True)
    shade(table.rows[-1].cells[0]); no_split(table.rows[-1])
    cells = table.add_row().cells
    set_text(cells[0], "Actor", bold=True)
    set_text(cells[1], "Hệ thống", bold=True)
    shade(cells[0], "EDF4FA"); shade(cells[1], "EDF4FA")
    no_split(table.rows[-1])
    for index, (actor_step, system_step) in enumerate(spec["flow"], start=1):
        cells = table.add_row().cells
        set_text(cells[0], f"{index}. {actor_step}")
        set_text(cells[1], f"{index}. {system_step}")
        no_split(table.rows[-1])
    row("Luồng thay thế", spec["alt"])
    row("Luồng ngoại lệ", spec["exc"])
    row("Đối chiếu hiện thực", spec["source"])
    repeat_header(table.rows[0])
    caption(doc, f"Bảng 3-{table_number} Đặc tả {spec['code']} – {spec['title']}", "table")


def add_class_summary(doc) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    for index, text in enumerate(("Sơ đồ", "Phạm vi", "Nhóm lớp/quan hệ trọng tâm")):
        set_text(table.rows[0].cells[index], text, True)
        shade(table.rows[0].cells[index])
    repeat_header(table.rows[0])
    data = [
        ("CLS-01", "Miền và dữ liệu", "User–Project–Task–Resource; Department, TaskGroup, RecurringTask, OptimizationResult, Notification, RefreshToken, ActivityLog, CompanySetting."),
        ("CLS-02", "Backend", "Route/controller, middleware quyền, bộ giải CSP–GA, nhập Excel, MongoDB/Mongoose, Socket.IO và nhật ký hoạt động."),
        ("CLS-03", "Frontend", "App, ProtectedRoute, context xác thực/theme/socket, các trang nghiệp vụ, API client và i18n."),
    ]
    for item in data:
        cells = table.add_row().cells
        for index, text in enumerate(item):
            set_text(cells[index], text)
        no_split(table.rows[-1])
    caption(doc, "Bảng 3-16 Phạm vi ba sơ đồ lớp của hệ thống", "table")


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite {OUTPUT}")
    required = ["UC00", *(f"UC{i:02d}" for i in range(1, 7)),
                *(f"ACT{i:02d}" for i in range(1, 11)),
                *(f"SEQ{i:02d}" for i in range(1, 11)),
                *(f"CLS{i:02d}" for i in range(1, 4))]
    missing = [code for code in required if not (IMAGES / f"{code}.png").exists()]
    if missing:
        raise FileNotFoundError(f"Missing UML images: {missing}")

    doc = Document(SOURCE)
    anchor = next(p for p in doc.paragraphs if p.style.name == "Heading 1" and "HIỆN THỰC" in p.text)
    original_elements = set(doc._element.body)
    figure_no = [3]

    doc.add_heading("Sơ đồ use case và phạm vi chức năng", level=2)
    paragraph(doc, "Phần này kế thừa cách trình bày của báo cáo mẫu Nhom03.pdf: sơ đồ use case tổng quát, danh mục chức năng, bảng đặc tả theo tác nhân–hệ thống, sau đó là activity, sequence và class diagram. Nội dung nghiệp vụ và sơ đồ được đối chiếu với mã nguồn ResourceAllocation; không sử dụng dữ liệu của hệ thống bán điện thoại trong mẫu.")
    paragraph(doc, "Sơ đồ tổng hợp UC-00 có một khung hệ thống chứa 72 use case và 6 tác nhân ở ngoài. Vì mật độ thông tin cao, các sơ đồ UC-01 đến UC-06 được kèm theo để đọc rõ từng phân hệ; các phần tử gốc có thể chỉnh sửa trong file Visual Paradigm ResourceAllocation_UML_MotKhung.vpp.")
    add_figure(doc, "UC00", "Sơ đồ use case tổng hợp của ResourceAllocation", figure_no)
    for index, title in enumerate((
        "Tổng quan hệ thống", "Xác thực, tài khoản và phân quyền",
        "Dự án, nhóm và công việc", "Nhân sự, phòng ban và lịch nghỉ",
        "Tối ưu hóa và benchmark", "Lịch, báo cáo, thông báo và cấu hình",
    ), start=1):
        add_figure(doc, f"UC{index:02d}", f"Sơ đồ use case phân hệ {title}", figure_no)

    doc.add_heading("Danh mục use case", level=2)
    paragraph(doc, "Mã UC-01.xx đến UC-06.xx được gán theo thứ tự phần tử của từng sơ đồ chi tiết. UC-01 là tầng tổng quan, do đó một số chức năng khái quát được triển khai thành các use case cụ thể ở UC-02 đến UC-06.")
    add_catalog(doc)

    doc.add_heading("Đặc tả các luồng nghiệp vụ trọng tâm", level=2)
    paragraph(doc, "Theo mẫu báo cáo, mỗi đặc tả nêu mục tiêu, tác nhân, tiền/hậu điều kiện, luồng chính, luồng thay thế và ngoại lệ. Mười một luồng dưới đây bao phủ các chuỗi xử lý trọng tâm; danh mục 72 use case phía trên bao quát toàn bộ chức năng ở mức sơ đồ. Mỗi hình activity/sequence được đặt ngay sau đặc tả liên quan khi có sơ đồ tương ứng.")
    for index, spec in enumerate(SPECS, start=1):
        heading = doc.add_heading(f"{spec['code']} – {spec['title']}", level=3)
        heading.paragraph_format.page_break_before = True
        add_spec_table(doc, spec, 4 + index)
        if spec["act"]:
            code = spec["act"]
            add_figure(doc, code, f"Sơ đồ activity {code} – {ACTIVITY_TITLES[code]}", figure_no)
        if spec["seq"]:
            code = spec["seq"]
            add_figure(doc, code, f"Sơ đồ sequence {code} – {SEQUENCE_TITLES[code]}", figure_no)

    doc.add_heading("Sơ đồ lớp và đối chiếu hiện thực", level=2)
    paragraph(doc, "Ba góc nhìn lớp được tách để người đọc phân biệt mô hình dữ liệu với cấu trúc backend và frontend. Sơ đồ mô hình miền cho thấy quan hệ dữ liệu; sơ đồ backend cho thấy đường đi qua điều khiển, phân quyền, dịch vụ và bộ giải; sơ đồ frontend cho thấy các thành phần giao diện và tầng gọi API.")
    add_class_summary(doc)
    for code, title in CLASS_TITLES.items():
        add_figure(doc, code, f"Sơ đồ class {code} – {title}", figure_no)
    paragraph(doc, "Các sơ đồ trong chương này được xuất từ dự án Visual Paradigm 10.0 ResourceAllocation_UML_MotKhung.vpp. Khi đọc bản điện tử có thể phóng to hình để kiểm tra thuộc tính, thông điệp và quan hệ; file .vpp đi kèm cho phép chỉnh sửa trực tiếp từng phần tử.")

    added = [element for element in doc._element.body if element not in original_elements]
    for element in added:
        anchor._element.addprevious(element)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.core_properties.subject = "ResourceAllocation – báo cáo đồ án với đặc tả use case và sơ đồ UML"
    doc.save(OUTPUT)
    print(OUTPUT)
    print(f"Added {figure_no[0] - 3} figures, {len(SPECS)} detailed specifications, 72-case catalog")


if __name__ == "__main__":
    main()
